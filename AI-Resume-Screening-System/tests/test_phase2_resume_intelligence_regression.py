# ============================================================
#  HireAI / TalentSync — Phase 2 Resume Intelligence & UX Regression
#  Covers GAP-01, GAP-03, GAP-04, GAP-08, GAP-09
# ============================================================

import os
import io
import sys
import time
import json
import unittest

_APP_ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), '..'))
if _APP_ROOT not in sys.path:
    sys.path.insert(0, _APP_ROOT)

import docx as python_docx
from reportlab.pdfgen import canvas  # type: ignore[import-untyped]
from werkzeug.datastructures import FileStorage

from app import create_app
from app.config.settings import TestingConfig
from app.database.connection import get_db
from app.controllers.auth_controller import (
    register_user,
    generate_password_reset_token,
    generate_email_verification_token
)
from app.controllers.resume_controller import process_resume_upload


def _make_docx_stream(text: str = "Senior Engineer with Python, Flask, SQL experience.") -> io.BytesIO:
    doc = python_docx.Document()
    doc.add_paragraph(text)
    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf


def _make_pdf_stream(text: str = "Candidate PDF Resume with Python and Docker.") -> io.BytesIO:
    buf = io.BytesIO()
    c = canvas.Canvas(buf)
    c.drawString(100, 750, text)
    c.save()
    buf.seek(0)
    return buf


class TestPhase2ResumeIntelligenceRegression(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.app = create_app(TestingConfig)

    def setUp(self):
        self.app_context = self.app.app_context()
        self.app_context.push()
        self.client = self.app.test_client()

    def tearDown(self):
        self.app_context.pop()

    # ------------------------------------------------------------
    # GAP-01: Resume Intelligence Data Contract
    # ------------------------------------------------------------
    def test_gap01_upload_response_includes_canonical_intel(self):
        """process_resume_upload must include canonical 'intel' dictionary in response data."""
        email = f"p2_cand_{int(time.time()*1000)}@test.com"
        reg = register_user("Intel Candidate", email, "ValidPassword10!", "candidate")
        self.assertTrue(reg["success"])
        user_id = reg["user"]["id"]

        stream = _make_docx_stream("Jorgen Sjoberg\nEmail: jorgen@example.com\nPython, Docker, SQL Developer")
        fs = FileStorage(stream=stream, filename="Jorgen_CV.docx", content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document")
        result = process_resume_upload(fs, user_id=user_id)

        self.assertTrue(result.get("success"))
        self.assertIn("data", result)
        data = result["data"]

        # Assert backward compatibility
        self.assertIn("ats_score", data)
        self.assertIn("skills", data)
        self.assertIsInstance(data["skills"], list)

        # Assert GAP-01: Canonical Intelligence object present
        self.assertIn("intel", data, "Upload response data must contain 'intel' object.")
        intel = data["intel"]
        self.assertIsInstance(intel, dict)
        self.assertIn("candidate", intel)
        self.assertIn("contact", intel)
        self.assertIn("skills", intel)
        self.assertIn("quality", intel)

    def test_gap01_structured_json_included_in_my_resumes(self):
        """GET /api/resume/my_resumes must return structured_json for each uploaded resume."""
        email = f"p2_cand_hist_{int(time.time()*1000)}@test.com"
        reg = register_user("History Candidate", email, "ValidPassword10!", "candidate")
        user_id = reg["user"]["id"]

        stream = _make_docx_stream("Python Developer with Cloud skills.")
        fs = FileStorage(stream=stream, filename="Candidate_Resume.docx", content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document")
        process_resume_upload(fs, user_id=user_id)

        with self.client.session_transaction() as sess:
            sess["user_id"] = user_id

        res = self.client.get("/api/resume/my_resumes")
        self.assertEqual(res.status_code, 200)
        resumes = res.get_json().get("resumes", [])
        self.assertGreater(len(resumes), 0)
        latest = resumes[0]
        self.assertIn("structured_json", latest)
        self.assertTrue(len(latest["structured_json"]) > 0)

    # ------------------------------------------------------------
    # GAP-03: Resume History Synchronization
    # ------------------------------------------------------------
    def test_gap03_unauthenticated_my_resumes_rejected(self):
        """Unauthenticated call to /api/resume/my_resumes must return 401."""
        res = self.client.get("/api/resume/my_resumes")
        self.assertEqual(res.status_code, 401)

    def test_gap03_user_resumes_isolated_per_account(self):
        """User A cannot see User B's resumes via /api/resume/my_resumes."""
        email_a = f"p2_user_a_{int(time.time()*1000)}@test.com"
        user_a = register_user("User A", email_a, "ValidPassword10!", "candidate")["user"]["id"]
        stream_a = _make_docx_stream("User A resume text")
        fs_a = FileStorage(stream=stream_a, filename="UserA.docx", content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document")
        process_resume_upload(fs_a, user_id=user_a)

        email_b = f"p2_user_b_{int(time.time()*1000)}@test.com"
        user_b = register_user("User B", email_b, "ValidPassword10!", "candidate")["user"]["id"]

        # Authenticate as User B
        with self.client.session_transaction() as sess:
            sess["user_id"] = user_b

        res = self.client.get("/api/resume/my_resumes")
        self.assertEqual(res.status_code, 200)
        resumes_b = res.get_json().get("resumes", [])
        # User B should have 0 resumes
        self.assertEqual(len(resumes_b), 0)

    # ------------------------------------------------------------
    # GAP-04: Resume Download & Preview Authorization
    # ------------------------------------------------------------
    def test_gap04_candidate_can_download_own_resume(self):
        """Candidate can download their own uploaded resume."""
        email = f"p2_dl_own_{int(time.time()*1000)}@test.com"
        user_id = register_user("DL User", email, "ValidPassword10!", "candidate")["user"]["id"]

        stream = _make_docx_stream("Downloadable resume content.")
        fs = FileStorage(stream=stream, filename="MyCV.docx", content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document")
        up = process_resume_upload(fs, user_id=user_id)
        resume_id = up["resume_id"]

        with self.client.session_transaction() as sess:
            sess["user_id"] = user_id

        res = self.client.get(f"/api/resume/download/{resume_id}")
        self.assertEqual(res.status_code, 200)
        self.assertIn("attachment", res.headers.get("Content-Disposition", ""))

    def test_gap04_candidate_preview_pdf_inline(self):
        """preview=1 on a PDF serves the file with inline content-disposition."""
        email = f"p2_prev_{int(time.time()*1000)}@test.com"
        user_id = register_user("Preview User", email, "ValidPassword10!", "candidate")["user"]["id"]

        stream = _make_pdf_stream("Preview PDF content.")
        fs = FileStorage(stream=stream, filename="PreviewCV.pdf", content_type="application/pdf")
        up = process_resume_upload(fs, user_id=user_id)
        resume_id = up["resume_id"]

        with self.client.session_transaction() as sess:
            sess["user_id"] = user_id

        res = self.client.get(f"/api/resume/download/{resume_id}?preview=1")
        self.assertEqual(res.status_code, 200)
        # Inline display (not attachment)
        self.assertNotIn("attachment", res.headers.get("Content-Disposition", ""))

    def test_gap04_candidate_cannot_download_other_candidate_resume(self):
        """Candidate B attempting to download Candidate A's resume receives 403 Forbidden."""
        user_a = register_user("Victim", f"vic_{int(time.time()*1000)}@test.com", "ValidPassword10!", "candidate")["user"]["id"]
        fs_a = FileStorage(stream=_make_docx_stream("Secret"), filename="Secret.docx", content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document")
        res_id_a = process_resume_upload(fs_a, user_id=user_a)["resume_id"]

        user_b = register_user("Attacker", f"atk_{int(time.time()*1000)}@test.com", "ValidPassword10!", "candidate")["user"]["id"]

        with self.client.session_transaction() as sess:
            sess["user_id"] = user_b

        res = self.client.get(f"/api/resume/download/{res_id_a}")
        self.assertEqual(res.status_code, 403)

    def test_gap04_recruiter_can_download_candidate_resume(self):
        """HR / Admin user is authorized to download candidate resumes."""
        user_c = register_user("Candidate C", f"cand_c_{int(time.time()*1000)}@test.com", "ValidPassword10!", "candidate")["user"]["id"]
        fs_c = FileStorage(stream=_make_docx_stream("Resume"), filename="ResumeC.docx", content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document")
        res_id = process_resume_upload(fs_c, user_id=user_c)["resume_id"]

        hr_user = register_user("HR Recruiter", f"hr_rec_{int(time.time()*1000)}@test.com", "ValidPassword10!", "hr", allow_privileged=True, is_verified=1)["user"]["id"]

        with self.client.session_transaction() as sess:
            sess["user_id"] = hr_user

        res = self.client.get(f"/api/resume/download/{res_id}")
        self.assertEqual(res.status_code, 200)

    # ------------------------------------------------------------
    # GAP-08: Password Reset & Email Verification Backend Flow
    # ------------------------------------------------------------
    def test_gap08_forgot_password_generates_token(self):
        """POST /api/auth/forgot_password generates token for valid email."""
        email = f"forgot_{int(time.time()*1000)}@test.com"
        register_user("Forgot User", email, "ValidPassword10!", "candidate")

        res = self.client.post("/api/auth/forgot_password", json={"email": email})
        self.assertEqual(res.status_code, 200)
        data = res.get_json()
        self.assertTrue(data.get("success"))
        self.assertIn("dev_reset_token", data)

    def test_gap08_reset_password_flow(self):
        """POST /api/auth/reset_password with valid token updates password."""
        email = f"reset_{int(time.time()*1000)}@test.com"
        register_user("Reset User", email, "OldPassword10!", "candidate", is_verified=1)
        tok_res = generate_password_reset_token(email)
        token = tok_res["dev_reset_token"]

        # Attempt reset with short password (< 10)
        res_short = self.client.post("/api/auth/reset_password", json={"token": token, "new_password": "Short9!"})
        self.assertEqual(res_short.status_code, 400)

        # Valid reset (>= 10)
        res_valid = self.client.post("/api/auth/reset_password", json={"token": token, "new_password": "NewValidPass10!"})
        self.assertEqual(res_valid.status_code, 200)
        self.assertTrue(res_valid.get_json().get("success"))

        # Verify login works with new password
        login_res = self.client.post("/api/auth/login", json={"email": email, "password": "NewValidPass10!"})
        self.assertEqual(login_res.status_code, 200)

    def test_gap08_email_verification_flow(self):
        """POST /api/auth/verify_email verifies unverified user token."""
        email = f"verify_{int(time.time()*1000)}@test.com"
        reg = register_user("Unverified User", email, "ValidPassword10!", "candidate")
        user_id = reg["user"]["id"]
        token = generate_email_verification_token(user_id)

        res = self.client.post("/api/auth/verify_email", json={"token": token})
        self.assertEqual(res.status_code, 200)
        self.assertTrue(res.get_json().get("success"))

        # Second verification should report already verified
        res_again = self.client.post("/api/auth/verify_email", json={"token": token})
        self.assertEqual(res_again.status_code, 400)

    # ------------------------------------------------------------
    # GAP-09: Recruiter Candidate View with Real Data
    # ------------------------------------------------------------
    def test_gap09_admin_candidates_returns_real_fields(self):
        """GET /api/admin/candidates returns degree, exp, and resume metadata without placeholders."""
        hr_email = f"hr_cand_{int(time.time()*1000)}@test.com"
        hr_id = register_user("HR Manager", hr_email, "ValidPassword10!", "hr", allow_privileged=True, is_verified=1)["user"]["id"]

        # Create a candidate who uploads a resume and applies to a job
        cand_email = f"cand_app_{int(time.time()*1000)}@test.com"
        cand_id = register_user("Real Candidate", cand_email, "ValidPassword10!", "candidate")["user"]["id"]
        fs = FileStorage(stream=_make_docx_stream("Software Engineer with 3 years experience at Acme Corp. B.Tech Computer Science."),
                         filename="RealResume.docx",
                         content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document")
        up = process_resume_upload(fs, user_id=cand_id)
        resume_id = up["resume_id"]

        with get_db() as conn:
            job = conn.execute("SELECT id FROM jobs LIMIT 1").fetchone()
            job_id = job["id"]
            conn.execute("INSERT INTO applications (user_id, job_id, match_score, status) VALUES (?, ?, ?, 'Reviewing')",
                         (cand_id, job_id, 88.5))
            conn.commit()

        with self.client.session_transaction() as sess:
            sess["user_id"] = hr_id

        res = self.client.get("/api/admin/candidates")
        self.assertEqual(res.status_code, 200)
        candidates = res.get_json()
        self.assertGreater(len(candidates), 0)

        found_cand = next((c for c in candidates if c.get("email") == cand_email), None)
        self.assertIsNotNone(found_cand)
        self.assertEqual(found_cand["resume_id"], resume_id)
        self.assertEqual(found_cand["resume_name"], "RealResume.docx")
        # Assert placeholders are NOT used
        self.assertNotEqual(found_cand["degree"], "See Profile")
        self.assertIn("degree", found_cand)
        self.assertIn("exp", found_cand)


if __name__ == "__main__":
    unittest.main()
