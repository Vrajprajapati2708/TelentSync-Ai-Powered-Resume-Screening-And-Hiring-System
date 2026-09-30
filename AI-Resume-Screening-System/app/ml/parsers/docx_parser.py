import io
import uuid
import math
import time
from datetime import datetime
import docx
from app.utils.logger import get_logger

logger = get_logger(__name__)

# Parser Configuration
PARSER_CONFIG = {
    "max_pages": 100,
    "max_file_size": 10 * 1024 * 1024,
    "encoding": "utf-8",
    "parser_version": "v1.0"
}

# Parser Status Constants
class ParseStatus:
    SUCCESS = "SUCCESS"
    FAILED = "FAILED"
    PARTIAL = "PARTIAL"
    UNSUPPORTED = "UNSUPPORTED"


def extract_docx_with_metadata(file, filename: str = "") -> dict:
    """
    Extract text from Word DOCX paragraphs and tables with metadata computation.
    """
    t_start = time.time()
    clean_name = filename or getattr(file, 'filename', '') or 'document.docx'

    logger.info(f"START PARSE [DOCX]: filename='{clean_name}'")

    if file is None:
        logger.warning("extract_docx_with_metadata received None file object.")
        duration_ms = int((time.time() - t_start) * 1000)
        return _build_docx_result(False, "", clean_name, 0, 0, 0, 0, ["File object is None."], duration_ms, ParseStatus.FAILED)

    # Calculate stream size safely
    file_size = 0
    try:
        if hasattr(file, 'seek') and hasattr(file, 'tell'):
            file.seek(0, 2)
            file_size = file.tell()
            file.seek(0)
    except Exception:
        file_size = 0

    try:
        if hasattr(file, 'seek'):
            file.seek(0)

        doc = docx.Document(file)
        lines = []

        # 1. Extract paragraph text
        for para in doc.paragraphs:
            txt = para.text.strip()
            if txt:
                lines.append(txt)

        # 2. Extract table cell text
        for table in doc.tables:
            for row in table.rows:
                row_cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if row_cells:
                    lines.append(" | ".join(row_cells))

        full_text = "\n".join(lines).strip()
        word_count = len(full_text.split())
        char_count = len(full_text)
        page_count = max(1, math.ceil(word_count / 450)) if word_count > 0 else 0
        duration_ms = int((time.time() - t_start) * 1000)

        if not full_text:
            msg = "DOCX file contains no readable text."
            logger.warning(msg)
            return _build_docx_result(False, "", clean_name, file_size, 0, 0, 0, [msg], duration_ms, ParseStatus.FAILED)

        logger.info(f"PARSED [DOCX]: status={ParseStatus.SUCCESS} | duration={duration_ms}ms | pages=~{page_count} | words={word_count}")
        return _build_docx_result(True, full_text, clean_name, file_size, page_count, word_count, char_count, [], duration_ms, ParseStatus.SUCCESS)

    except Exception as e:
        duration_ms = int((time.time() - t_start) * 1000)
        err_msg = f"Failed to parse DOCX: {str(e)}"
        logger.error(f"PARSED FAILED [DOCX]: duration={duration_ms}ms | error={err_msg}")
        return _build_docx_result(False, "", clean_name, file_size, 0, 0, 0, [err_msg], duration_ms, ParseStatus.FAILED)


def extract_text_from_docx(file) -> str:
    """
    Backward-compatible wrapper: extracts text from DOCX file.
    """
    res = extract_docx_with_metadata(file)
    return res['raw_text'] if res['success'] else ""


def _build_docx_result(success: bool, raw_text: str, filename: str, size: int, pages: int, words: int, chars: int, errors: list, duration_ms: int = 0, status: str = ParseStatus.SUCCESS) -> dict:
    """Helper to construct standardized DOCX parser result structure."""
    return {
        'success': success,
        'raw_text': raw_text,
        'metadata': {
            'document_id': f"doc_{uuid.uuid4().hex[:12]}",
            'file_name': filename,
            'file_type': 'docx',
            'encoding': PARSER_CONFIG["encoding"],
            'file_size_bytes': size,
            'page_count': pages,
            'word_count': words,
            'character_count': chars,
            'language': 'en',
            'parser_name': 'python-docx',
            'parser_version': PARSER_CONFIG["parser_version"],
            'parser_duration_ms': duration_ms,
            'processing_time_ms': duration_ms,
            'created_at': datetime.now().isoformat(),
            'parse_status': status,
            'parse_errors': errors
        },
        'errors': errors
    }



